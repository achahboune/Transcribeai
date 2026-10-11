"""Plan-gated post-processing helpers for TranscribeAI.

These helpers deliberately keep the transcript as the source of truth. The
subtitle exporters use Whisper segment timestamps, while the speaker labels
are explicitly marked as estimated turns unless a dedicated diarization
provider is configured later.
"""
from __future__ import annotations

import io
import re
from typing import Any

from .config import PLAN_LIMITS_MINUTES


# "creator" and "pro" are legacy plans: they now have the same access as "unlimited".
PLAN_RANK = {"free": 0, "unlimited": 1, "creator": 1, "pro": 1}


def has_plan(plan: str, minimum: str) -> bool:
    return PLAN_RANK.get(plan, 0) >= PLAN_RANK.get(minimum, 99)


def clean_dictionary(value: str | None) -> list[str]:
    if not value:
        return []
    terms = []
    for raw in re.split(r"[,\n]", value):
        term = " ".join(raw.strip().split())
        if term and term not in terms:
            terms.append(term[:80])
    return terms[:50]


def build_transcription_prompt(dictionary: list[str]) -> str | None:
    if not dictionary:
        return None
    return (
        "Use these custom names, brands, and technical terms when they are "
        "spoken. Preserve the original language and do not add terms that "
        "were not spoken: " + ", ".join(dictionary)
    )


def _timestamp(seconds: float) -> str:
    seconds = max(0.0, float(seconds or 0))
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    whole = int(seconds % 60)
    millis = int(round((seconds - int(seconds)) * 1000))
    if millis == 1000:
        whole += 1
        millis = 0
    return f"{hours:02d}:{minutes:02d}:{whole:02d},{millis:03d}"


def _vtt_timestamp(seconds: float) -> str:
    return _timestamp(seconds).replace(",", ".")


def normalize_segments(segments: list[dict[str, Any]] | None, transcript: str) -> list[dict[str, Any]]:
    if segments:
        return [
            {
                "start": float(item.get("start") or 0),
                "end": float(item.get("end") or item.get("start") or 0),
                "text": str(item.get("text") or "").strip(),
            }
            for item in segments
            if str(item.get("text") or "").strip()
        ]
    return [{"start": 0.0, "end": 0.0, "text": transcript.strip()}] if transcript.strip() else []


def transcript_to_srt(segments: list[dict[str, Any]], diarized: bool = False) -> str:
    lines = []
    for idx, segment in enumerate(segments, 1):
        text = segment["text"]
        if diarized:
            text = f"[Speaker {((idx - 1) % 2) + 1} — estimated] {text}"
        end = segment["end"] if segment["end"] > segment["start"] else segment["start"] + 2
        lines += [str(idx), f"{_timestamp(segment['start'])} --> {_timestamp(end)}", text, ""]
    return "\n".join(lines)


def transcript_to_vtt(segments: list[dict[str, Any]], diarized: bool = False) -> str:
    lines = ["WEBVTT", ""]
    for segment in segments:
        text = segment["text"]
        if diarized:
            text = f"[Speaker {((len(lines) - 2) % 2) + 1} — estimated] {text}"
        end = segment["end"] if segment["end"] > segment["start"] else segment["start"] + 2
        lines += [f"{_vtt_timestamp(segment['start'])} --> {_vtt_timestamp(end)}", text, ""]
    return "\n".join(lines)


def transcript_to_docx(transcript: str, summary: str | None = None, title: str = "TranscribeAI transcript") -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    if summary:
        document.add_heading("Summary", level=2)
        document.add_paragraph(summary)
    document.add_heading("Transcript", level=2)
    for paragraph in transcript.splitlines() or [transcript]:
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def transcript_to_pdf(transcript: str, summary: str | None = None, title: str = "TranscribeAI transcript") -> bytes:
    """Build a simple PDF. Uses DejaVu Sans (Unicode) with text shaping so that
    accents and Arabic render correctly; falls back to a core font if the
    font file is not available."""
    from fpdf import FPDF

    font_path = next(
        (p for p in (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        ) if __import__("os").path.exists(p)),
        None,
    )
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    unicode_ok = False
    if font_path:
        try:
            pdf.add_font("DejaVu", "", font_path)
            pdf.set_font("DejaVu", size=11)
            try:
                pdf.set_text_shaping(True)
            except Exception:
                pass
            unicode_ok = True
        except Exception:
            unicode_ok = False
    if not unicode_ok:
        pdf.set_font("Helvetica", size=11)

    def safe(text: str) -> str:
        return text if unicode_ok else text.encode("latin-1", "replace").decode("latin-1")

    def write(text: str, size: int = 11, gap: float = 6):
        pdf.set_font_size(size)
        for line in (text.splitlines() or [text]):
            rtl = any("\u0590" <= ch <= "\u08ff" for ch in line)
            pdf.multi_cell(0, gap, safe(line) or " ", align="R" if rtl else "L",
                           new_x="LMARGIN", new_y="NEXT")

    write(title, size=16, gap=9)
    pdf.ln(3)
    if summary:
        write("Summary", size=13, gap=8)
        write(summary)
        pdf.ln(3)
    write("Transcript", size=13, gap=8)
    write(transcript)
    return bytes(pdf.output())


def plan_capabilities(plan: str) -> dict[str, bool]:
    return {
        "summary": has_plan(plan, "creator"),
        "docx_export": has_plan(plan, "creator"),
        "pdf_export": has_plan(plan, "creator"),
        "srt_export": has_plan(plan, "creator"),
        "vtt_export": has_plan(plan, "creator"),
        "custom_dictionary": has_plan(plan, "creator"),
        "estimated_speaker_turns": has_plan(plan, "creator"),
        "advanced_subtitles": has_plan(plan, "creator"),
        "team": has_plan(plan, "creator"),
        "api_access": has_plan(plan, "creator"),
    }
